"""Offline regression tests for document boundaries and lossless export."""
import os
import tempfile
import unittest
import zipfile
from unittest.mock import patch

import pymupdf

from bs4 import BeautifulSoup
from ebooklib import epub
import docx

from core.parser import BookParser, BookProject, BookChapter, BookParagraph
from core.exporter import BookExporter, is_academic_paper


class DocumentWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def test_markdown_subheadings_do_not_create_chapters(self):
        path = self.path('novel.md')
        with open(path, 'w', encoding='utf-8') as stream:
            stream.write('# Chapter One\n\nFirst scene.\n\n## A memory\n\nSecond scene.\n\n# Chapter Two\n\nLast scene.')
        project = BookParser.parse_file(path, 'markdown', document_type='novel')
        self.assertEqual([c.title for c in project.chapters], ['Chapter One', 'Chapter Two'])
        self.assertIn('A memory', [p.original_text for p in project.chapters[0].paragraphs])
        self.assertEqual(project.document_type, 'novel')

    def test_prologue_roman_chapter_and_epilogue(self):
        path = self.path('story.txt')
        with open(path, 'w', encoding='utf-8') as stream:
            stream.write('Prologue\n\nAn opening.\n\nChapter IV\n\nA story.\n\nEpilogue\n\nAn ending.')
        project = BookParser.parse_file(path, 'plain')
        self.assertEqual([c.title for c in project.chapters], ['Prologue', 'Chapter IV', 'Epilogue'])

    def test_docx_first_heading_names_first_chapter(self):
        path = self.path('story.docx')
        document = docx.Document()
        document.add_heading('Prologue', level=1)
        document.add_paragraph('An opening.')
        document.add_heading('A memory', level=2)
        document.add_paragraph('Still the opening.')
        document.add_heading('Chapter I', level=1)
        document.add_paragraph('Main story.')
        document.save(path)
        project = BookParser.parse_file(path, 'word')
        self.assertEqual([c.title for c in project.chapters], ['Prologue', 'Chapter I'])

    def test_docx_strongest_heading_level_and_body_reference(self):
        path = self.path('heading-two.docx')
        document = docx.Document()
        document.add_heading('Opening', level=2)
        document.add_paragraph('Chapter 2 explains why this happens.')
        document.add_heading('A detail', level=3)
        document.add_heading('Ending', level=2)
        document.add_paragraph('The end.')
        document.save(path)
        project = BookParser.parse_file(path, 'heading-two')
        self.assertEqual([c.title for c in project.chapters], ['Opening', 'Ending'])
        self.assertEqual([p.tag for p in project.chapters[0].paragraphs], ['h2', 'p', 'h3'])

    def test_docx_plain_roman_chapters_do_not_split_body_references(self):
        path = self.path('plain-headings.docx')
        document = docx.Document()
        for text in ('Prologue', 'An opening.', 'Chapter IV', 'Chapter 4 explains the next event.', 'Epilogue', 'The end.'):
            document.add_paragraph(text)
        document.save(path)
        project = BookParser.parse_file(path, 'plain-headings')
        self.assertEqual([c.title for c in project.chapters], ['Prologue', 'Chapter IV', 'Epilogue'])
        self.assertIn('Chapter 4 explains the next event.', [p.original_text for p in project.chapters[1].paragraphs])

    def test_invalid_workflow_and_empty_document(self):
        path = self.path('empty.txt')
        with open(path, 'w', encoding='utf-8') as stream:
            stream.write('')
        for mode in ('paper', 'invalid', 'novel'):
            with self.assertRaises(ValueError):
                BookParser.parse_file(path, 'bad', document_type=mode)

    def test_explicit_type_controls_export_style(self):
        project = BookProject('test', 'Story', chapters=[BookChapter('chap_0', 'Abstract')], document_type='novel')
        self.assertFalse(is_academic_paper(project))
        project.document_type = 'paper'
        project.chapters[0].title = 'Methods'
        self.assertTrue(is_academic_paper(project))

    def test_pdf_sections_keep_subsections_and_top_page_heading(self):
        path = self.path('paper.pdf')
        with pymupdf.open() as document:
            page = document.new_page()
            page.insert_text((60, 50), 'A Synthetic Paper', fontsize=20, fontname='hebo')
            page.insert_text((60, 110), 'Abstract', fontsize=14, fontname='hebo')
            page.insert_text((60, 140), 'This is an abstract describing the synthetic experiment.', fontsize=10)
            page.insert_text((60, 190), '1 Introduction', fontsize=14, fontname='hebo')
            page.insert_text((60, 220), 'We investigate structure preservation in synthetic papers.', fontsize=10)
            page = document.new_page()
            page.insert_text((60, 50), '2 Methods', fontsize=14, fontname='hebo')
            page.insert_text((60, 90), 'This method uses a simple fixture for regression testing.', fontsize=10)
            page.insert_text((60, 140), '2.1 Results', fontsize=12, fontname='hebo')
            page.insert_text((60, 180), '1 Configure settings', fontsize=10)
            page.insert_text((60, 210), 'A regular numbered list should not become a new section.', fontsize=10)
            document.save(path)
        with patch('core.parser.os.makedirs'):
            project = BookParser.parse_file(path, 'synthetic', document_type='paper')
        self.assertEqual([c.title for c in project.chapters],
                         ['Phần mở đầu / Tiêu đề', 'Abstract', '1 Introduction', '2 Methods'])
        methods = project.chapters[-1]
        self.assertIn(('2.1 Results', 'h3'), [(p.original_text, p.tag) for p in methods.paragraphs])
        self.assertTrue(any('1 Configure settings' in p.original_text for p in methods.paragraphs))

    def test_pdf_two_column_text_order(self):
        path = self.path('columns.pdf')
        with pymupdf.open() as document:
            page = document.new_page()
            for x, y, text in [(330, 160, 'Right bottom.'), (60, 160, 'Left bottom.'),
                               (330, 100, 'Right top.'), (60, 100, 'Left top.')]:
                page.insert_text((x, y), text, fontsize=10)
            document.save(path)
        with patch('core.parser.os.makedirs'):
            project = BookParser._parse_pdf_mupdf(path, 'columns')
        text = [p.original_text for c in project.chapters for p in c.paragraphs]
        self.assertEqual(text, ['Left top.', 'Left bottom.', 'Right top.', 'Right bottom.'])

    def test_epub_mixed_containers_preserve_all_text_once(self):
        from core.epub_structure import epub_text_elements, _extract_element_text
        soup = BeautifulSoup('<div>Before <b>bold</b><h1>Chapter One</h1><ul><li>Intro<p>Child</p>After</li></ul>End</div>', 'html.parser')
        elements = epub_text_elements(soup)
        self.assertEqual([_extract_element_text(el) for el in elements],
                         ['Before bold', 'Chapter One', 'Intro', 'Child', 'After', 'End'])
        self.assertEqual([_extract_element_text(el) for el in epub_text_elements(soup)],
                         ['Before bold', 'Chapter One', 'Intro', 'Child', 'After', 'End'])

    def test_epub_mixed_part_grouping_and_leaf_toc(self):
        book = epub.EpubBook()
        book.set_identifier('mixed-parts')
        book.set_title('Mixed Parts')
        book.set_language('en')
        items = []
        for name, title in [('prologue', 'Prologue'), ('one', 'Chapter One'), ('two', 'Chapter Two')]:
            item = epub.EpubHtml(title=title, file_name=f'{name}.xhtml')
            item.content = f'<h1>{title}</h1><p>Text for {name}.</p>'
            book.add_item(item)
            items.append(item)
        book.toc = (epub.Link('prologue.xhtml', 'Prologue', 'prologue'),
                    (epub.Section('Part I', 'one.xhtml'), [
                        epub.Link('one.xhtml', 'Chapter One', 'one'),
                        epub.Link('two.xhtml', 'Chapter Two', 'two')]))
        book.spine = items
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        path = self.path('mixed-parts.epub')
        epub.write_epub(path, book)
        project = BookParser.parse_file(path, 'mixed-parts')
        self.assertEqual([c.title for c in project.chapters], ['Prologue', 'Chapter One', 'Chapter Two'])

    def test_epub_inline_text_and_line_breaks(self):
        from core.epub_structure import _extract_element_text
        soup = BeautifulSoup('<p>Hello <b>world</b> today<br/>Another line.</p>', 'html.parser')
        self.assertEqual(_extract_element_text(soup.p), 'Hello world today Another line.')

    def make_epub(self):
        book = epub.EpubBook()
        book.set_identifier('test-locator')
        book.set_title('Two Chapters')
        book.set_language('en')
        chapter = epub.EpubHtml(title='Story', file_name='text/story.xhtml')
        chapter.content = '<h1 id="one">Chapter One</h1><p>Repeated line.</p><h1 id="two">Chapter Two</h1><p>Repeated line.</p>'
        book.add_item(chapter)
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        book.toc = (epub.Link('text/story.xhtml#one', 'Chapter One', 'one'), epub.Link('text/story.xhtml#two', 'Chapter Two', 'two'))
        book.spine = ['nav', chapter]
        path = self.path('source.epub')
        epub.write_epub(path, book)
        return path

    def test_split_epub_export_updates_all_chapters_and_duplicates(self):
        source = self.make_epub()
        project = BookParser.parse_file(source, 'epub')
        self.assertEqual(len(project.chapters), 2)
        for index, chapter in enumerate(project.chapters):
            for paragraph in chapter.paragraphs:
                paragraph.translated_text = f'{index}:{paragraph.original_text}'
                paragraph.status = 'done'
        for bilingual in (False, True):
            output = self.path(f'export-{bilingual}.epub')
            BookExporter.export_epub(project, output, bilingual=bilingual)
            with zipfile.ZipFile(output) as archive:
                name = next(n for n in archive.namelist() if n.endswith('/text/story.xhtml'))
                soup = BeautifulSoup(archive.read(name), 'html.parser')
                text = soup.get_text(' ', strip=True)
                self.assertIn('0:Repeated line.', text)
                self.assertIn('1:Repeated line.', text)
                self.assertIn('1:Chapter Two', text)
                self.assertEqual(archive.infolist()[0].filename, 'mimetype')
                self.assertEqual(archive.infolist()[0].compress_type, zipfile.ZIP_STORED)

    def test_legacy_epub_duplicates_consumed_in_source_order(self):
        source = self.make_epub()
        project = BookProject('legacy', 'Old story', source_file_path=source, source_format='epub', chapters=[
            BookChapter('chap_0', 'One', doc_name='text/story.xhtml', paragraphs=[
                BookParagraph('p0', 'Repeated line.', translated_text='First translation.'),
                BookParagraph('p1', 'Repeated line.', translated_text='Second translation.'),
            ])
        ])
        output = self.path('legacy.epub')
        BookExporter.export_epub(project, output)
        with zipfile.ZipFile(output) as archive:
            name = next(n for n in archive.namelist() if n.endswith('/text/story.xhtml'))
            text = BeautifulSoup(archive.read(name), 'html.parser').get_text(' ', strip=True)
            self.assertIn('First translation.', text)
            self.assertIn('Second translation.', text)


if __name__ == '__main__':
    unittest.main()
