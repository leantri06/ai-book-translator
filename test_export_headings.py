"""Offline export checks for translatable headings and document metadata."""
import os
import tempfile
import unittest
import zipfile

from bs4 import BeautifulSoup
import docx
from ebooklib import epub

from core.exporter import BookExporter
from core.parser import BookChapter, BookParagraph, BookProject


class ExportHeadingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.project = BookProject('headings', 'Research', document_type='paper', chapters=[
            BookChapter('chap_0', '1 Introduction', paragraphs=[
                BookParagraph('p0', '1 Introduction', translated_text='1 Giới thiệu', tag='h2'),
                BookParagraph('p1', '1.1 Context', translated_text='1.1 Bối cảnh', tag='h3'),
                BookParagraph('p2', 'Details', translated_text='Chi tiết', tag='h5'),
                BookParagraph('p3', 'Body text.', translated_text='Nội dung.'),
            ])
        ])

    def output(self, name):
        return os.path.join(self.tmp.name, name)

    def test_html_heading_not_duplicated(self):
        for bilingual in (False, True):
            path = self.output(f'{bilingual}.html')
            BookExporter.export_html(self.project, path, bilingual)
            with open(path, encoding='utf-8') as stream:
                soup = BeautifulSoup(stream.read(), 'html.parser')
            chapter = soup.select_one('.chapter-section')
            self.assertEqual(len(chapter.find_all('h2')), 1)
            self.assertEqual(chapter.h2.get_text(), '1 Giới thiệu')
            self.assertEqual(chapter.h3.get_text(), '1.1 Bối cảnh')
            self.assertEqual(chapter.h5.get_text(), 'Chi tiết')
            self.assertEqual(chapter.get_text().count('1 Introduction'), int(bilingual))

    def test_fresh_epub_heading_not_duplicated(self):
        for bilingual in (False, True):
            path = self.output(f'{bilingual}.epub')
            BookExporter.export_epub(self.project, path, bilingual)
            with zipfile.ZipFile(path) as archive:
                name = next(n for n in archive.namelist() if n.endswith('/chap_0.xhtml'))
                soup = BeautifulSoup(archive.read(name), 'html.parser')
            self.assertEqual(len(soup.find_all('h2')), 1)
            self.assertEqual(soup.h2.get_text(), '1 Giới thiệu')
            self.assertEqual(soup.h3.get_text(), '1.1 Bối cảnh')
            self.assertEqual(soup.h5.get_text(), 'Chi tiết')
            self.assertEqual(soup.body.get_text().count('1 Introduction'), int(bilingual))

    def test_docx_retains_heading_styles(self):
        for bilingual in (False, True):
            path = self.output(f'{bilingual}.docx')
            BookExporter.export_docx(self.project, path, bilingual)
            document = docx.Document(path)
            headings = [p for p in document.paragraphs if p.style.name.startswith('Heading')]
            self.assertEqual([p.style.name for p in headings], ['Heading 2', 'Heading 3', 'Heading 5'])
            self.assertEqual(headings[0].text, '1 Introduction\n1 Giới thiệu' if bilingual else '1 Giới thiệu')

    def test_txt_heading_not_duplicated(self):
        for bilingual in (False, True):
            path = self.output(f'{bilingual}.txt')
            BookExporter.export_txt(self.project, path, bilingual)
            with open(path, encoding='utf-8') as stream:
                text = stream.read()
            self.assertNotIn('--- 1 Introduction ---', text)
            self.assertEqual(text.count('1 Giới thiệu'), 1)
            self.assertEqual(text.count('1 Introduction'), int(bilingual))

    def test_body_matching_title_still_gets_generated_heading(self):
        self.project.chapters[0].paragraphs[0].tag = 'p'
        path = self.output('body.html')
        BookExporter.export_html(self.project, path)
        with open(path, encoding='utf-8') as stream:
            chapter = BeautifulSoup(stream.read(), 'html.parser').select_one('.chapter-section')
        self.assertEqual(chapter.h2.get_text(), '1 Introduction')
        self.assertEqual(chapter.p.get_text(), '1 Giới thiệu')

    def test_untranslated_book_is_not_vietnamese_original(self):
        for chapter in self.project.chapters:
            for paragraph in chapter.paragraphs:
                paragraph.translated_text = ''
        path = self.output('untranslated.epub')
        BookExporter.export_epub(self.project, path)
        book = epub.read_epub(path)
        self.assertIn('Bản Dịch Tiếng Việt', book.get_metadata('DC', 'title')[0][0])

    def test_explicit_textbook_preserves_original_title(self):
        self.project.document_type = 'textbook'
        path = self.output('textbook.epub')
        BookExporter.export_epub(self.project, path)
        book = epub.read_epub(path)
        self.assertEqual(book.get_metadata('DC', 'title')[0][0], 'Research')


if __name__ == '__main__':
    unittest.main()
